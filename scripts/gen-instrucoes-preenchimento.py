#!/usr/bin/env python3
# Gera as Instruções de Preenchimento (Versões A e B) da MARIAH a partir da spec canônica.
# Fonte única: spec/mariah-spec.json. Saída: .docx limpos (v2), estilo institucional MARIAH.
# Uso: gen-instrucoes-preenchimento.py [spec] [outdir] [locale]  (locale: pt-BR padrão | es)
# pt-BR = versão normativa (arquivo canônico). Outro locale: textos da matriz vêm dos campos
# i18n da spec (aprovados pelo Z); a moldura vem de STR[locale]; saída com sufixo -<locale>.
import json, sys, os, re
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SPEC = sys.argv[1] if len(sys.argv) > 1 else 'spec/mariah-spec.json'
OUTDIR = sys.argv[2] if len(sys.argv) > 2 else '/tmp/out'
LOCALE = sys.argv[3] if len(sys.argv) > 3 else 'pt-BR'
os.makedirs(OUTDIR, exist_ok=True)
spec = json.load(open(SPEC, encoding='utf-8'))

def L(node, field):
    """Campo traduzido (spec i18n) ou canônico pt-BR — mesma regra do label() da UI."""
    v = (node.get('i18n') or {}).get(LOCALE, {}).get(field)
    return v if isinstance(v, str) else node.get(field)

STR = {
 'pt-BR': {
  'hdr': ['Cenário','Máx.','Nível I','Nível II','Nível III','Nível IV'],
  'faixa': '{a} a {b}',
  'semBanco': 'Sem banco de dados', 'comBanco': 'Com banco (Bloco 6.b)',
  'fElim': 'ELIMINATÓRIA', 'fDilQuant': 'Diligência (0 pt; devolução)', 'fEvid': 'Evidência (só-abate)',
  'fMitig': 'Mitigação (subtrai)', 'fCond': 'Condicional', 'fNa': 'admite “não se aplica”', 'fDilQual': 'Diligência (devolução)',
  'sim': 'Sim', 'nao': 'Não',
  'colItem': 'Item', 'colPergunta': 'Pergunta e orientação para preenchimento',
  'respRisco': 'Resposta de risco: ', 'peso': 'Peso: ', 'pts': '{n} pts', 'orientacao': 'Orientação. ',
  'titulo': 'Instruções de Preenchimento',
  'matriz': 'Matriz de Avaliação de Risco de Inteligência Artificial em Pesquisa com Seres Humanos',
  'status': 'Documento gerado a partir da especificação canônica da matriz (v2). Versão aprovada pela INAEP em 16/09/2026 — em revisão editorial para publicação.',
  'cortesia': None,
  'ctxTitulo': 'Caracterização do contexto de uso (não pontua)',
  'ctxTexto': 'Antes da pontuação, o protocolo responde a oito perguntas descritivas obrigatórias. Não pontuam e não determinam o nível por si mesmas: delimitam o objeto da avaliação e modulam a leitura dos eixos/blocos.',
  'ctxCol': 'Pergunta descritiva',
  'bTitulo': 'Versão B — Quantitativa', 'bSub': 'Pontuação ponderada em sete blocos de avaliação',
  'bComo': '1. Como funciona a Versão B',
  'bP1': 'A Versão B é quantitativa: cada resposta de risco soma pontos, distribuídos em sete blocos temáticos. A pontuação total é a soma dos blocos (modelo aditivo), e os pontos de corte definem o nível final do protocolo. Cada questão tem um peso fixo, indicado na coluna “Item”.',
  'bP2': 'Mecanismos que alteram a aritmética simples: (a) a Cláusula de Prevalência Ética (Bloco 4) — se P4.1 ou P4.2 for “Sim”, o protocolo é forçado ao Nível IV; (b) a eliminatória de cadeia de custódia (P6.b.2), que torna o protocolo “não avaliável pela MARIAH”; (c) o Bloco 7 (Mitigação), bidirecional e dividido em três subblocos — 7A (consultas, somam quando ausentes), 7B (medidas de redução, subtraem quando presentes) e 7C (evidências de transparência, regime de só-abate: presentes subtraem até 23 pts, ausentes não somam); (d) o Bloco 6.b condicional (Res. CNS n.º 738/2024), que soma 29 pts ao Bloco 6 quando há banco de dados e integra três itens de diligência (P6.b.4.1, P6.b.6, P6.b.7) que não pontuam e, quando não atendidos, tornam o protocolo não avaliável no mérito (devolução); e (e) a diligência de novo consentimento (P2.8), eliminatória para sistemas adaptativos sem plano de novo consentimento. O piso do Bloco 7 é zero: a pontuação não fica negativa.',
  'passo0': 'Passo 0 — Filtro de banco de dados (Res. CNS n.º 738/2024)',
  'cortes': 'Pontos de corte',
  'cortesNota': ('Com banco de dados, o teto teórico é 304 e o teto avaliável é 297 (descontada a questão eliminatória P6.b.2). '
           'Os cortes derivam das frações fixas do baseline 238 (50/238, 110/238, 180/238) aplicadas ao teto teórico vigente, '
           'com arredondamento de meia-unidade para cima. A calibração definitiva é prospectiva.'),
  'bQuestoes': '2. Questões por bloco', 'maxBloco': '{nome} (máx. {n} pts)',
  'aTitulo': 'Versão A — Qualitativa', 'aSub': 'Percurso por cinco eixos (mais o Eixo 3.b condicional)',
  'aComo': '1. Como funciona a Versão A',
  'aP1': 'A Versão A é qualitativa: perguntas de resposta binária, organizadas em cinco eixos temáticos, percorridos em sequência. Em cada eixo, qualquer resposta de risco eleva o nível do eixo; o nível final do protocolo é o MÁXIMO entre os eixos — o risco só sobe, nunca desce.',
  'aP2': 'Mecanismos: (a) a regra especial do Eixo 3.b (banco de dados, Res. CNS n.º 738/2024): 0 respostas de risco → Nível I; 1–2 → Nível III; 3 ou mais → Nível IV; (b) a eliminatória de cadeia de custódia (3.b.2), que torna o protocolo “não avaliável pela MARIAH”; (c) os itens de diligência do Eixo 3.b, que não pontuam e, quando não atendidos, devolvem o protocolo ao pesquisador; e (d) a diligência de novo consentimento (2.10), eliminatória para sistemas adaptativos sem plano de novo consentimento. Na Versão A as mitigações não alteram o nível.',
  'aQuestoes': '2. Questões por eixo',
 },
 'es': {
  'hdr': ['Escenario','Máx.','Nivel I','Nivel II','Nivel III','Nivel IV'],
  'faixa': '{a} a {b}',
  'semBanco': 'Sin banco de datos', 'comBanco': 'Con banco (Bloque 6.b)',
  'fElim': 'ELIMINATORIA', 'fDilQuant': 'Diligencia (0 pt; devolución)', 'fEvid': 'Evidencia (solo descuenta)',
  'fMitig': 'Mitigación (resta)', 'fCond': 'Condicional', 'fNa': 'admite «No se aplica»', 'fDilQual': 'Diligencia (devolución)',
  'sim': 'Sí', 'nao': 'No',
  'colItem': 'Ítem', 'colPergunta': 'Pregunta y orientación para la cumplimentación',
  'respRisco': 'Respuesta de riesgo: ', 'peso': 'Peso: ', 'pts': '{n} pts', 'orientacao': 'Orientación. ',
  'titulo': 'Instrucciones de Cumplimentación',
  'matriz': 'Matriz de Evaluación de Riesgo de Inteligencia Artificial en Investigación con Seres Humanos',
  'status': 'Documento generado a partir de la especificación canónica de la matriz (v2). Versión aprobada por la INAEP el 16/09/2026 — en revisión editorial para su publicación.',
  'cortesia': 'Esta es una traducción de cortesía. La versión normativa vigente es la versión en portugués (pt-BR).',
  'ctxTitulo': 'Caracterización del contexto de uso (no puntúa)',
  'ctxTexto': 'Antes de la puntuación, el protocolo responde a ocho preguntas descriptivas obligatorias. No puntúan ni determinan el nivel por sí mismas: delimitan el objeto de la evaluación y modulan la lectura de los ejes/bloques.',
  'ctxCol': 'Pregunta descriptiva',
  'bTitulo': 'Versión B — Cuantitativa', 'bSub': 'Puntuación ponderada en siete bloques de evaluación',
  'bComo': '1. Cómo funciona la Versión B',
  'bP1': 'La Versión B es cuantitativa: cada respuesta de riesgo suma puntos, distribuidos en siete bloques temáticos. La puntuación total es la suma de los bloques (modelo aditivo), y los puntos de corte definen el nivel final del protocolo. Cada pregunta tiene un peso fijo, indicado en la columna «Ítem».',
  'bP2': 'Mecanismos que alteran la aritmética simple: (a) la Cláusula de Primacía Ética (Bloque 4) — si P4.1 o P4.2 es «Sí», el protocolo se eleva obligatoriamente al Nivel IV; (b) la eliminatoria de cadena de custodia (P6.b.2), que vuelve el protocolo «no evaluable por la MARIAH»; (c) el Bloque 7 (Mitigación), bidireccional y dividido en tres subbloques — 7A (consultas, suman cuando están ausentes), 7B (medidas de reducción, restan cuando están presentes) y 7C (evidencias de transparencia, régimen de solo descuento: presentes restan hasta 23 pts, ausentes no suman); (d) el Bloque 6.b condicional (Res. CNS n.º 738/2024), que suma 29 pts al Bloque 6 cuando hay banco de datos e integra tres ítems de diligencia (P6.b.4.1, P6.b.6, P6.b.7) que no puntúan y, cuando no se cumplen, vuelven el protocolo no evaluable en el mérito (devolución); y (e) la diligencia de nuevo consentimiento (P2.8), eliminatoria para sistemas adaptativos sin plan de nuevo consentimiento. El piso del Bloque 7 es cero: la puntuación no queda negativa.',
  'passo0': 'Paso 0 — Filtro de banco de datos (Res. CNS n.º 738/2024)',
  'cortes': 'Puntos de corte',
  'cortesNota': ('Con banco de datos, el techo teórico es 304 y el techo evaluable es 297 (descontada la pregunta eliminatoria P6.b.2). '
           'Los cortes derivan de las fracciones fijas de la línea de base 238 (50/238, 110/238, 180/238) aplicadas al techo teórico vigente, '
           'con redondeo de media unidad hacia arriba. La calibración definitiva es prospectiva.'),
  'bQuestoes': '2. Preguntas por bloque', 'maxBloco': '{nome} (máx. {n} pts)',
  'aTitulo': 'Versión A — Cualitativa', 'aSub': 'Recorrido por cinco ejes (más el Eje 3.b condicional)',
  'aComo': '1. Cómo funciona la Versión A',
  'aP1': 'La Versión A es cualitativa: preguntas de respuesta binaria, organizadas en cinco ejes temáticos, recorridos en secuencia. En cada eje, cualquier respuesta de riesgo eleva el nivel del eje; el nivel final del protocolo es el MÁXIMO entre los ejes — el riesgo solo sube, nunca baja.',
  'aP2': 'Mecanismos: (a) la regla especial del Eje 3.b (banco de datos, Res. CNS n.º 738/2024): 0 respuestas de riesgo → Nivel I; 1–2 → Nivel III; 3 o más → Nivel IV; (b) la eliminatoria de cadena de custodia (3.b.2), que vuelve el protocolo «no evaluable por la MARIAH»; (c) los ítems de diligencia del Eje 3.b, que no puntúan y, cuando no se cumplen, devuelven el protocolo al investigador; y (d) la diligencia de nuevo consentimiento (2.10), eliminatoria para sistemas adaptativos sin plan de nuevo consentimiento. En la Versión A las mitigaciones no alteran el nivel.',
  'aQuestoes': '2. Preguntas por eje',
 },
}[LOCALE]
SUFIXO = '' if LOCALE == 'pt-BR' else f'-{LOCALE}'

TEAL = RGBColor(0x0F, 0x76, 0x6E)
SHADE = "D5E8F0"
GREY = RGBColor(0x55, 0x55, 0x55)

def set_cell_shade(cell, color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd'); shd.set(qn('w:val'),'clear'); shd.set(qn('w:color'),'auto'); shd.set(qn('w:fill'),color)
    tcPr.append(shd)

def base_style(doc):
    st = doc.styles['Normal']; st.font.name='Arial'; st.font.size=Pt(10.5)
    st.element.rPr.rFonts.set(qn('w:eastAsia'),'Arial')
    for section in doc.sections:
        section.page_height=Cm(29.7); section.page_width=Cm(21.0)
        section.top_margin=Cm(2.2); section.bottom_margin=Cm(2.0)
        section.left_margin=Cm(2.2); section.right_margin=Cm(2.2)

def add_run(p, text, bold=False, color=None, size=None, italic=False):
    r=p.add_run(text); r.bold=bold; r.italic=italic
    if color is not None: r.font.color.rgb=color
    if size is not None: r.font.size=Pt(size)
    r.font.name='Arial'
    return r

def h1(doc, text):
    p=doc.add_paragraph(); p.space_after=Pt(4); add_run(p,text,bold=True,color=TEAL,size=15); return p
def h2(doc, text):
    p=doc.add_paragraph(); add_run(p,text,bold=True,color=TEAL,size=12.5); p.paragraph_format.space_before=Pt(10); p.paragraph_format.space_after=Pt(4); return p
def para(doc, text, size=10.5, color=None, italic=False):
    p=doc.add_paragraph(); add_run(p,text,size=size,color=color,italic=italic); p.paragraph_format.space_after=Pt(6); return p

def cutoffs_table(doc, tb, tc, notas):
    t=doc.add_table(rows=3, cols=6); t.style='Table Grid'; t.alignment=WD_TABLE_ALIGNMENT.CENTER
    hdr=STR['hdr']
    for j,htext in enumerate(hdr):
        c=t.rows[0].cells[j]; set_cell_shade(c,SHADE); p=c.paragraphs[0]; add_run(p,htext,bold=True,size=9.5)
    def row(i, cen, mx, faixas):
        cells=t.rows[i].cells
        vals=[cen,str(mx)]+faixas
        for j,v in enumerate(vals):
            p=cells[j].paragraphs[0]; add_run(p,v,size=9.5,bold=(j==0))
    def faixas(th):
        I,II,III=th['levelI'],th['levelII'],th['levelIII']
        f=STR['faixa']
        return [f.format(a=0,b=I), f.format(a=I+1,b=II), f.format(a=II+1,b=III), f'≥ {III+1}']
    row(1,STR['semBanco'], tb['maxScore'], faixas(tb))
    row(2,STR['comBanco'], tc['maxScore'], faixas(tc))
    para(doc, notas, size=8.5, color=GREY, italic=True)

def flags_quant(q):
    fl=[]
    ef=q.get('efeito')
    if q.get('eliminatorio') and q.get('pontos',0)!=0: fl.append(STR['fElim'])
    if ef=='diligencia': fl.append(STR['fDilQuant'])
    if ef=='evidencia': fl.append(STR['fEvid'])
    if ef=='mitigacao': fl.append(STR['fMitig'])
    if q.get('exibicaoCondicional'): fl.append(STR['fCond'])
    if q.get('hasNaOption'): fl.append(STR['fNa'])
    return fl

def flags_qual(q):
    fl=[]
    if q.get('eliminatorio'): fl.append(STR['fElim'])
    if q.get('naoPontuavel') or q.get('efeito')=='diligencia': fl.append(STR['fDilQual'])
    if q.get('exibicaoCondicional'): fl.append(STR['fCond'])
    if q.get('hasNaOption'): fl.append(STR['fNa'])
    return fl

def risk_label(ra):
    return {'sim':STR['sim'],'nao':STR['nao'],'na':'—'}.get((ra or '').lower(), ra or '—')

ID_DISPLAY = {'contexto1':'C.1','contexto2':'C.2'}
def disp_id(qid):
    return ID_DISPLAY.get(qid, qid)

def clean_dica(text):
    if not text: return text
    t=text
    # remove referências internas de engenharia (fichas F-##, carimbo de versão)
    t=re.sub(r'\s*[—-]\s*F-?\d+\b','',t)
    t=re.sub(r'\s*\(\s*F-?\d+\s*\)','',t)
    t=re.sub(r'\bF-?\d+\b','',t)
    t=re.sub(r'\s*\(\s*v\d+(?:\.\d+)*\s*\)','',t)
    t=re.sub(r'\s{2,}',' ',t)
    t=re.sub(r'\s+([,.;])',r'\1',t)
    t=re.sub(r'\(\s*\)','',t)
    return t.strip()

def question_table(doc, questoes, quant):
    t=doc.add_table(rows=1, cols=2); t.style='Table Grid'
    t.columns[0].width=Cm(4.6); t.columns[1].width=Cm(12.4)
    hc=t.rows[0].cells
    set_cell_shade(hc[0],SHADE); set_cell_shade(hc[1],SHADE)
    add_run(hc[0].paragraphs[0],STR['colItem'],bold=True,size=9.5)
    add_run(hc[1].paragraphs[0],STR['colPergunta'],bold=True,size=9.5)
    for q in questoes:
        row=t.add_row().cells
        # coluna Item
        p=row[0].paragraphs[0]; add_run(p,disp_id(q['id']),bold=True,size=9.5)
        p2=row[0].add_paragraph(); add_run(p2,STR['respRisco'],size=8.5); add_run(p2,risk_label(q.get('riskAnswer')),bold=True,size=8.5)
        if quant:
            pts=q.get('pontos',0)
            p3=row[0].add_paragraph(); add_run(p3,STR['peso'],size=8.5); add_run(p3,STR['pts'].format(n=(f'{pts:+d}' if pts else '0')),bold=True,size=8.5)
        for fl in (flags_quant(q) if quant else flags_qual(q)):
            pf=row[0].add_paragraph(); add_run(pf,fl,size=8,color=GREY,italic=True)
        # coluna Pergunta + orientação
        pp=row[1].paragraphs[0]; add_run(pp,L(q,'pergunta'),bold=True,size=9.5)
        if q.get('dica'):
            pd=row[1].add_paragraph(); add_run(pd,STR['orientacao'],bold=True,size=9); add_run(pd,clean_dica(L(q,'dica')),size=9)

def header_block(doc, versao_titulo, subtitulo):
    p=doc.add_paragraph(); add_run(p,'MARIAH',bold=True,color=TEAL,size=20); p.paragraph_format.space_after=Pt(0)
    p=doc.add_paragraph(); add_run(p,STR['titulo'],bold=True,size=13)
    p=doc.add_paragraph(); add_run(p,versao_titulo,bold=True,size=12,color=TEAL)
    para(doc, subtitulo, size=9.5, color=GREY)
    para(doc, STR['matriz'], size=9.5, color=GREY, italic=True)
    para(doc, STR['status'], size=8.5, color=GREY, italic=True)
    if STR['cortesia']: para(doc, STR['cortesia'], size=8.5, color=GREY, italic=True)

def context_section(doc, ctxs):
    h2(doc,STR['ctxTitulo'])
    para(doc,STR['ctxTexto'],size=9.5)
    t=doc.add_table(rows=1,cols=2); t.style='Table Grid'
    t.columns[0].width=Cm(2.2); t.columns[1].width=Cm(14.8)
    hc=t.rows[0].cells; set_cell_shade(hc[0],SHADE); set_cell_shade(hc[1],SHADE)
    add_run(hc[0].paragraphs[0],STR['colItem'],bold=True,size=9.5); add_run(hc[1].paragraphs[0],STR['ctxCol'],bold=True,size=9.5)
    for c in ctxs:
        row=t.add_row().cells
        add_run(row[0].paragraphs[0], disp_id(c.get('id','C')),bold=True,size=9.5)
        add_run(row[1].paragraphs[0], L(c,'pergunta'),size=9.5)
        if c.get('dica'):
            pd=row[1].add_paragraph(); add_run(pd,clean_dica(L(c,'dica')),size=8.5,color=GREY)

# ---------------- VERSÃO B ----------------
def gen_B():
    doc=Document(); base_style(doc)
    header_block(doc,STR['bTitulo'],STR['bSub'])
    h2(doc,STR['bComo'])
    para(doc,STR['bP1'])
    para(doc,STR['bP2'])
    h2(doc,STR['passo0'])
    dbf=spec['databaseFilterQuestion']
    para(doc, L(dbf,'pergunta'))
    para(doc, L(dbf,'dica') or '', size=9.5, color=GREY)
    context_section(doc, spec['contextQuestions'])
    h2(doc,STR['cortes'])
    notas=STR['cortesNota']
    cutoffs_table(doc, spec['thresholdsBase'], spec['thresholdsComBanco'], notas)
    h2(doc,STR['bQuestoes'])
    for b in spec['quantitativeBlocks']:
        h2(doc, STR['maxBloco'].format(nome=L(b,'nome'), n=b['maxPontos']))
        if b.get('descricao'): para(doc,L(b,'descricao'),size=9.5,color=GREY)
        question_table(doc, b['questoes'], quant=True)
    out=os.path.join(OUTDIR,f'instrucoes-preenchimento-versao-b-mariah{SUFIXO}.docx'); doc.save(out); return out

# ---------------- VERSÃO A ----------------
def gen_A():
    doc=Document(); base_style(doc)
    header_block(doc,STR['aTitulo'],STR['aSub'])
    h2(doc,STR['aComo'])
    para(doc,STR['aP1'])
    para(doc,STR['aP2'])
    h2(doc,STR['passo0'])
    dbf=spec['databaseFilterQuestion']
    para(doc, L(dbf,'pergunta')); para(doc, L(dbf,'dica') or '', size=9.5, color=GREY)
    context_section(doc, spec['contextQuestions'])
    h2(doc,STR['aQuestoes'])
    for a in spec['qualitativeAxes']:
        h2(doc, L(a,'nome'))
        if a.get('descricao'): para(doc,L(a,'descricao'),size=9.5,color=GREY)
        question_table(doc, a['questoes'], quant=False)
    out=os.path.join(OUTDIR,f'instrucoes-preenchimento-versao-a-mariah{SUFIXO}.docx'); doc.save(out); return out

print(gen_B())
print(gen_A())
